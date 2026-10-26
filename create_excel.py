import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import CellIsRule, DataBarRule, ColorScaleRule
from openpyxl.chart import BarChart, Reference, PieChart, LineChart

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Dataset"

# Headers
headers = [
    "Student ID", "Student Name", "Maths", "Science", "English", "Computer",
    "Total Marks", "Average", "Highest Marks", "Lowest Marks", "Percentage", "Grade", "Result"
]
ws.append(headers)

# Data
data = [
    ["S101", "Aarav", 85, 78, 90, 88],
    ["S102", "Diya", 72, 81, 76, 80],
    ["S103", "Kabir", 95, 89, 92, 97],
    ["S104", "Mehak", 65, 70, 68, 75],
    ["S105", "Arjun", 88, 84, 86, 91],
    ["S106", "Simran", 91, 94, 89, 93],
    ["S107", "Rohan", 56, 62, 60, 58],
    ["S108", "Ananya", 79, 83, 81, 85],
    ["S109", "Yash", 67, 72, 70, 74],
    ["S110", "Priya", 98, 96, 95, 99],
]

for row in data:
    ws.append(row)

# Add formulas
for i in range(2, 12):
    ws[f"G{i}"] = f"=SUM(C{i}:F{i})"
    ws[f"H{i}"] = f"=ROUND(AVERAGE(C{i}:F{i}), 2)"
    ws[f"I{i}"] = f"=MAX(C{i}:F{i})"
    ws[f"J{i}"] = f"=MIN(C{i}:F{i})"
    ws[f"K{i}"] = f"=ROUND((G{i}/400)*100, 2)"
    grade_formula = (
        f'=IF(K{i}>=90, "A+", '
        f'IF(K{i}>=80, "A", '
        f'IF(K{i}>=70, "B", '
        f'IF(K{i}>=60, "C", '
        f'IF(K{i}>=50, "D", "F")))))'
    )
    ws[f"L{i}"] = grade_formula
    ws[f"M{i}"] = f'=IF(K{i}>=40, "Pass", "Fail")'

# Conditional Formatting
green_fill = PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")
red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
lightblue_fill = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")

marks_range = "C2:F11"
ws.conditional_formatting.add(marks_range, CellIsRule(operator='greaterThanOrEqual', formula=['90'], fill=green_fill))
ws.conditional_formatting.add(marks_range, CellIsRule(operator='lessThan', formula=['60'], fill=red_fill))

# Data Bars
rule_databar = DataBarRule(start_type='min', end_type='max', color="638EC6", showValue=None, minLength=None, maxLength=None)
ws.conditional_formatting.add("K2:K11", rule_databar)

# Color Scale
rule_colorscale = ColorScaleRule(start_type='min', start_color='F8696B',
                                 mid_type='percentile', mid_value=50, mid_color='FFEB84',
                                 end_type='max', end_color='63BE7B')
ws.conditional_formatting.add("G2:G11", rule_colorscale)

# Highlight A+
ws.conditional_formatting.add("L2:L11", CellIsRule(operator='equal', formula=['"A+"'], fill=lightblue_fill))

# Averages for Chart 2
ws["A13"] = "Average"
ws["C13"] = "=AVERAGE(C2:C11)"
ws["D13"] = "=AVERAGE(D2:D11)"
ws["E13"] = "=AVERAGE(E2:E11)"
ws["F13"] = "=AVERAGE(F2:F11)"

# Charts
ws_chart1 = wb.create_sheet(title="Chart 1")
chart1 = BarChart()
chart1.type = "col"
chart1.style = 10
chart1.title = "Student-wise Total Marks"
chart1.x_axis.title = "Student Name"
chart1.y_axis.title = "Total Marks"
data1 = Reference(ws, min_col=7, min_row=1, max_row=11)
cats1 = Reference(ws, min_col=2, min_row=2, max_row=11)
chart1.add_data(data1, titles_from_data=True)
chart1.set_categories(cats1)
ws_chart1.add_chart(chart1, "A1")

ws_chart2 = wb.create_sheet(title="Chart 2")
chart2 = PieChart()
chart2.title = "Subject-wise Average Marks"
data2 = Reference(ws, min_col=3, max_col=6, min_row=13, max_row=13)
cats2 = Reference(ws, min_col=3, max_col=6, min_row=1, max_row=1)
chart2.add_data(data2, from_rows=True, titles_from_data=False)
chart2.set_categories(cats2)
ws_chart2.add_chart(chart2, "A1")

ws_chart3 = wb.create_sheet(title="Chart 3")
chart3 = LineChart()
chart3.title = "Percentage Trend of Students"
chart3.x_axis.title = "Student Name"
chart3.y_axis.title = "Percentage"
data3 = Reference(ws, min_col=11, min_row=1, max_row=11)
cats3 = Reference(ws, min_col=2, min_row=2, max_row=11)
chart3.add_data(data3, titles_from_data=True)
chart3.set_categories(cats3)
ws_chart3.add_chart(chart3, "A1")

wb.save(r"E:\Assignment_Solution.xlsx")
print("Excel sheet created successfully at E:\Assignment_Solution.xlsx")
